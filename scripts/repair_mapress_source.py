"""Apply the user's 2026-10-05 confirmed source/metadata corrections."""
from pathlib import Path
from collections import Counter
import hashlib
import openpyxl

root = Path(__file__).resolve().parents[1]
p = root / 'RawSource/Family_Upload_vs_Annotation_Tickets_Checked.xlsx'
w = openpyxl.load_workbook(p)
s = w['Family_vs_Tickets']
h = {c.value: c.column for c in s[1]}
r = next(r for r in range(2, s.max_row + 1) if s.cell(r, 1).value == '434_PF_CO_cCap_Mapress')
assert str(s.cell(r, h['Ticket_IDs']).value) == '72579'
w['Meta']['B15'] = s.max_row - 1
w['Meta']['B31'] = 0
w.save(p)
p = root / 'RawSource/Annotation_RFA_equivalence_Checked.xlsx'
w = openpyxl.load_workbook(p)
s = w['Equivalence']
counts = Counter(s.cell(r, 6).value for r in range(2, s.max_row + 1))
w['Meta']['B5'] = s.max_row - 1
for r in range(2, 5):
    w['Summary'].cell(r, 2, counts[w['Summary'].cell(r, 1).value])
w['Summary']['B5'] = s.max_row - 1
w.save(p)
v = p.read_bytes()
sha = hashlib.sha1(b'blob ' + str(len(v)).encode() + b'\0' + v).hexdigest()
for name in ['scripts/build_dashboard_data.py', 'tests/equivalence.test.js']:
    q = root / name
    q.write_text(q.read_text(encoding='utf-8').replace('01a9ea0a39b5c1d106b45a6d3778162b40614e49', sha), encoding='utf-8')
print(sha)
